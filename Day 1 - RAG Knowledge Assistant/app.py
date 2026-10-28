import io
import re
import time
import hashlib
from datetime import datetime

import faiss
import numpy as np
import streamlit as st
from sentence_transformers import SentenceTransformer
from pypdf import PdfReader
from docx import Document


# ============================================================
# PAGE CONFIGURATION
# ============================================================
st.set_page_config(
    page_title="Nexus | HR Knowledge Assistant",
    page_icon="🧠",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ============================================================
# STYLING
# ============================================================
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Space+Grotesk:wght@500;600;700&display=swap');

:root {
    --bg: #0b1020;
    --panel: #121a2e;
    --border: #26334d;
    --text: #edf3ff;
    --muted: #91a1bc;
    --accent: #8b9cff;
    --cyan: #5eead4;
}

.stApp {
    background:
        radial-gradient(ellipse at 10% 0%, #182748 0%, transparent 45%),
        linear-gradient(135deg, #090e1b, #10182b);
    color: var(--text);
    font-family: 'DM Sans', sans-serif;
}

h1, h2, h3 {
    font-family: 'Space Grotesk', sans-serif !important;
    color: var(--text);
}

.block-container {
    padding-top: 2rem;
    padding-bottom: 3rem;
    max-width: 1450px;
}

[data-testid="stSidebar"] {
    background: rgba(12, 19, 35, 0.97);
    border-right: 1px solid var(--border);
}

[data-testid="stSidebar"] * {
    color: var(--text);
}

.hero {
    padding: 27px;
    border: 1px solid #344466;
    border-radius: 22px;
    background: linear-gradient(120deg, #172746, #19213a 60%, #142d3b);
    margin-bottom: 22px;
}

.eyebrow {
    color: var(--cyan);
    font-size: 0.75rem;
    font-weight: 700;
    letter-spacing: 2px;
    text-transform: uppercase;
}

.hero h1 {
    margin: 10px 0 8px 0;
    font-size: clamp(2rem, 4vw, 3.1rem);
    letter-spacing: -1.5px;
}

.hero p {
    color: #bac8e2;
    margin-bottom: 0;
    font-size: 1rem;
}

.metric {
    padding: 17px;
    background: rgba(19, 29, 49, 0.92);
    border: 1px solid var(--border);
    border-radius: 15px;
    min-height: 112px;
}

.metric-label {
    color: var(--muted);
    font-size: 0.83rem;
}

.metric-value {
    font-size: 1.7rem;
    font-weight: 700;
    margin-top: 8px;
    color: #eaf0ff;
}

.section-card {
    padding: 19px;
    background: rgba(18, 26, 46, 0.84);
    border: 1px solid var(--border);
    border-radius: 17px;
}

.stButton > button {
    border-radius: 10px;
    border: 1px solid #3b4d72;
    min-height: 42px;
    font-weight: 600;
    transition: 0.2s ease;
}

.stButton > button:hover {
    border-color: #a5b4fc;
    color: #c7d2fe;
}

.stChatMessage {
    background: rgba(22, 32, 53, 0.8);
    border: 1px solid #293650;
    border-radius: 15px;
    padding: 13px;
}

[data-testid="stChatInput"] {
    border-color: #354563;
}

[data-testid="stFileUploader"] {
    background: rgba(20, 30, 50, 0.6);
    border: 1px dashed #445576;
    border-radius: 13px;
    padding: 10px;
}

.small-muted {
    color: #91a1bc;
    font-size: 0.85rem;
}

div[data-testid="stExpander"] {
    border: 1px solid var(--border);
    border-radius: 12px;
}

hr {
    border-color: #293650;
}

footer {
    visibility: hidden;
}
</style>
""", unsafe_allow_html=True)


# ============================================================
# DEMO DATA
# Replace or supplement this data with your actual HR corpus.
# Uploaded documents are added to this demo corpus.
# ============================================================
DEMO_POLICIES = {
    "HR Policy — Annual Leave": """
Annual Leave Policy

Employees are entitled to 24 days of annual leave per calendar year.
Leave requests should normally be submitted at least 5 working days
in advance and require manager approval. Employees should check their
official employment terms for eligibility and carry-forward rules.
""",
    "HR Policy — Sick Leave": """
Sick Leave Policy

Employees who cannot work due to illness should inform their manager
before the start of the working day. A medical certificate may be
required for extended periods of absence. Employees should follow the
official HR procedure for reporting and recording sick leave.
""",
    "HR Policy — Paternity Leave": """
Paternity Leave Policy

Eligible employees are entitled to 15 working days of paternity leave.
The leave should be taken within six months of the child's birth.
Employees should contact HR to confirm eligibility and documentation
requirements before submitting their leave request.
""",
    "HR Policy — Remote Work": """
Work From Home Policy

Employees may work from home up to 2 days per week with manager
approval. Employees must remain available during normal working hours
and comply with company information security requirements.
""",
    "HR Policy — Travel Expenses": """
Business Travel and Expense Policy

Employees travelling for business can claim reasonable hotel,
transportation, and meal expenses in accordance with company policy.
Original receipts should be submitted within 30 days of returning
from the trip. Expenses remain subject to the applicable approval
and reimbursement procedures.
""",
}


# ============================================================
# SESSION STATE
# ============================================================
if "messages" not in st.session_state:
    st.session_state.messages = []

if "documents" not in st.session_state:
    st.session_state.documents = {}

if "last_sources" not in st.session_state:
    st.session_state.last_sources = []

if "search_count" not in st.session_state:
    st.session_state.search_count = 0

if "last_question" not in st.session_state:
    st.session_state.last_question = ""


# ============================================================
# DOCUMENT EXTRACTION
# ============================================================
def extract_document(uploaded_file):
    name = uploaded_file.name
    raw = uploaded_file.getvalue()
    extension = name.lower().rsplit(".", 1)[-1]

    if extension == "txt":
        return raw.decode("utf-8", errors="replace")

    if extension == "pdf":
        reader = PdfReader(io.BytesIO(raw))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return "\n".join(pages)

    if extension == "docx":
        document = Document(io.BytesIO(raw))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))

        return "\n".join(paragraphs)

    raise ValueError("Supported formats are PDF, DOCX and TXT.")


def chunk_text(text, chunk_size=900, overlap=150):
    """Split long documents into overlapping text chunks."""
    text = re.sub(r"\s+", " ", text).strip()

    if not text:
        return []

    chunks = []
    start = 0

    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        if end == len(text):
            break

        start = end - overlap

    return chunks


# ============================================================
# EMBEDDING MODEL AND FAISS INDEX
# ============================================================
@st.cache_resource(show_spinner=False)
def load_model():
    return SentenceTransformer("sentence-transformers/all-MiniLM-L6-v2")


@st.cache_data(show_spinner=False, max_entries=3)
def build_index(corpus_signature, source_texts):
    """
    source_texts is a tuple of (source_name, text) tuples.
    The signature invalidates the cached index when documents change.
    """
    del corpus_signature

    passages = []

    for source_name, text in source_texts:
        for chunk_number, chunk in enumerate(chunk_text(text), start=1):
            passages.append({
                "source": source_name,
                "chunk": chunk_number,
                "text": chunk,
            })

    if not passages:
        return None, []

    model = load_model()

    vectors = model.encode(
        [p["text"] for p in passages],
        convert_to_numpy=True,
        normalize_embeddings=True,
        show_progress_bar=False,
    )

    vectors = np.asarray(vectors, dtype="float32")
    index = faiss.IndexFlatIP(vectors.shape[1])
    index.add(vectors)

    return index, passages


def get_corpus():
    corpus = dict(DEMO_POLICIES)
    corpus.update(st.session_state.documents)

    source_texts = tuple(sorted(corpus.items()))

    signature = hashlib.sha256(
        repr(source_texts).encode("utf-8")
    ).hexdigest()

    return source_texts, signature


def retrieve(question, top_k=5, min_score=0.20):
    source_texts, signature = get_corpus()
    index, passages = build_index(signature, source_texts)

    if index is None or not passages:
        return []

    model = load_model()
    query_vector = model.encode(
        [question],
        convert_to_numpy=True,
        normalize_embeddings=True,
    )
    query_vector = np.asarray(query_vector, dtype="float32")

    count = min(top_k, len(passages))
    scores, indices = index.search(query_vector, count)

    results = []

    for score, idx in zip(scores[0], indices[0]):
        if idx < 0:
            continue

        score = float(score)

        if score < min_score:
            continue

        result = dict(passages[idx])
        result["score"] = score
        results.append(result)

    return results


# ============================================================
# ANSWER PRESENTATION
# This is extractive retrieval, NOT LLM-generated reasoning.
# ============================================================
def make_response(question, results):
    if not results:
        return (
            "I couldn't find a sufficiently relevant passage in the "
            "currently indexed documents. Try rephrasing your question "
            "or upload the policy document containing the answer."
        )

    best = results[0]

    return (
        "I found relevant information in the indexed policy documents. "
        "The highest-ranked passage is shown below. Please check the "
        "source text to verify the exact wording and applicable conditions."
        "\n\n**Most relevant passage**\n\n"
        + best["text"]
    )


def render_sources(sources):
    if not sources:
        st.info("No relevant source passages were found.")
        return

    for i, source in enumerate(sources, start=1):
        score_pct = max(0, min(100, round(source["score"] * 100)))

        with st.expander(
            f"📄 {source['source']}  ·  Match score {score_pct}%"
        ):
            st.write(source["text"])
            st.caption(
                f"Chunk {source['chunk']} · "
                f"Cosine similarity: {source['score']:.3f}"
            )


# ============================================================
# SIDEBAR
# ============================================================
with st.sidebar:
    st.markdown("## 🧠 Nexus")
    st.caption("HR KNOWLEDGE ASSISTANT")
    st.divider()

    page = st.radio(
        "Navigation",
        ["💬 Assistant", "📚 Knowledge Base", "⚙️ Settings"],
        label_visibility="collapsed",
    )

    st.divider()
    st.markdown("### System status")
    st.success("Streamlit interface active")
    st.info("FAISS semantic retrieval")

    corpus, signature = get_corpus()
    source_count = len(corpus)

    st.metric("Indexed source documents", source_count)
    st.metric("Conversation messages", len(st.session_state.messages))

    st.divider()
    st.caption("Model: all-MiniLM-L6-v2")
    st.caption("Index: FAISS inner product / cosine similarity")


# ============================================================
# HEADER AND METRICS
# ============================================================
st.markdown("""
<div class="hero">
    <div class="eyebrow">Enterprise knowledge · Retrieval augmented search</div>
    <h1>Find answers in your HR policies.</h1>
    <p>
        Search indexed documents, inspect source passages, and explore
        your knowledge base through a conversational interface.
    </p>
</div>
""", unsafe_allow_html=True)

corpus, signature = get_corpus()

m1, m2, m3, m4 = st.columns(4)

with m1:
    st.markdown(
        f'<div class="metric"><div class="metric-label">Source documents</div>'
        f'<div class="metric-value">{len(corpus)}</div></div>',
        unsafe_allow_html=True,
    )

with m2:
    st.markdown(
        f'<div class="metric"><div class="metric-label">User questions</div>'
        f'<div class="metric-value">{st.session_state.search_count}</div></div>',
        unsafe_allow_html=True,
    )

with m3:
    st.markdown(
        f'<div class="metric"><div class="metric-label">Chat messages</div>'
        f'<div class="metric-value">{len(st.session_state.messages)}</div></div>',
        unsafe_allow_html=True,
    )

with m4:
    st.markdown(
        '<div class="metric"><div class="metric-label">Retrieval engine</div>'
        '<div class="metric-value">FAISS</div></div>',
        unsafe_allow_html=True,
    )

st.write("")


# ============================================================
# ASSISTANT PAGE
# ============================================================
if page == "💬 Assistant":

    left, right = st.columns([1.8, 1], gap="large")

    with left:
        st.subheader("💬 Conversation")
        st.caption("Ask a question about the currently indexed documents.")

        suggestions = [
            "How much annual leave is available?",
            "What is the paternity leave policy?",
            "What are the work-from-home rules?",
            "How do I claim travel expenses?",
        ]

        st.markdown("**Try a question**")

        suggestion_cols = st.columns(2)

        for i, suggestion in enumerate(suggestions):
            with suggestion_cols[i % 2]:
                if st.button(
                    suggestion,
                    key=f"suggestion_{i}",
                    use_container_width=True,
                ):
                    st.session_state.pending_question = suggestion

        st.divider()

        for i, message in enumerate(st.session_state.messages):
            with st.chat_message(message["role"]):
                st.markdown(message["content"])

                if message.get("sources"):
                    render_sources(message["sources"])

        typed_question = st.chat_input(
            "Ask about annual leave, remote work, expenses..."
        )

        question = typed_question or st.session_state.pop(
            "pending_question", None
        )

        if question:
            st.session_state.messages.append({
                "role": "user",
                "content": question,
            })
            st.session_state.search_count += 1
            st.session_state.last_question = question

            with st.spinner("Searching your knowledge base..."):
                try:
                    results = retrieve(
                        question,
                        top_k=5,
                        min_score=0.20,
                    )
                    answer = make_response(question, results)

                except Exception as exc:
                    results = []
                    answer = (
                        "An error occurred while searching the knowledge "
                        "base. Check the application logs and ensure the "
                        "embedding model loaded correctly.\n\n"
                        f"Error: `{str(exc)[:300]}`"
                    )

            st.session_state.messages.append({
                "role": "assistant",
                "content": answer,
                "sources": results,
            })

            st.session_state.last_sources = results
            st.rerun()

    with right:
        st.subheader("🔎 Search insights")

        if st.session_state.last_question:
            st.markdown("**Latest query**")
            st.write(st.session_state.last_question)
        else:
            st.info("Your latest query and retrieval details will appear here.")

        st.markdown("**Latest retrieved passages**")

        if st.session_state.last_sources:
            for result in st.session_state.last_sources[:3]:
                score = result["score"]
                st.markdown(f"**{result['source']}**")
                st.progress(max(0.0, min(1.0, score)))
                st.caption(f"Similarity: {score:.3f}")
        else:
            st.caption("No search results yet.")

        st.divider()
        st.markdown("**Actions**")

        if st.button("🗑️ Clear chat", use_container_width=True):
            st.session_state.messages = []
            st.session_state.last_sources = []
            st.session_state.last_question = ""
            st.rerun()

        if st.session_state.messages:
            export_text = []

            for message in st.session_state.messages:
                export_text.append(
                    f"{message['role'].upper()}\n{message['content']}"
                )

                for source in message.get("sources", []):
                    export_text.append(
                        f"Source: {source['source']}\n"
                        f"Similarity: {source['score']:.3f}\n"
                        f"{source['text']}"
                    )

            st.download_button(
                "⬇️ Export conversation",
                data="\n\n".join(export_text),
                file_name="nexus_conversation.txt",
                mime="text/plain",
                use_container_width=True,
            )


# ============================================================
# KNOWLEDGE BASE PAGE
# ============================================================
elif page == "📚 Knowledge Base":

    st.subheader("📚 Knowledge Base")
    st.write(
        "Upload company policies to add them to the searchable corpus. "
        "Supported formats: PDF, DOCX and TXT."
    )

    st.info(
        "The built-in HR policies are demonstration data. "
        "Upload your actual policy documents to search your own content."
    )

    uploaded_files = st.file_uploader(
        "Upload policy documents",
        type=["pdf", "docx", "txt"],
        accept_multiple_files=True,
        help="Text-based PDFs work best. Scanned PDFs require OCR.",
    )

    if st.button("➕ Process uploaded documents", type="primary"):
        if not uploaded_files:
            st.warning("Choose at least one document first.")
        else:
            added = 0

            with st.spinner("Extracting text from documents..."):
                for uploaded in uploaded_files:
                    try:
                        extracted = extract_document(uploaded)

                        if not extracted.strip():
                            st.warning(
                                f"No extractable text found in {uploaded.name}."
                            )
                            continue

                        st.session_state.documents[uploaded.name] = extracted
                        added += 1

                    except Exception as exc:
                        st.error(
                            f"Could not process {uploaded.name}: {exc}"
                        )

            if added:
                st.success(f"Processed {added} document(s).")
                st.cache_data.clear()
                st.rerun()

    st.divider()
    st.markdown("### Indexed source documents")

    combined = dict(DEMO_POLICIES)
    combined.update(st.session_state.documents)

    for name, text in combined.items():
        with st.expander(f"📄 {name}"):
            st.write(text[:5000])
            if len(text) > 5000:
                st.caption("Preview truncated.")
            st.caption(f"{len(text):,} characters")

    uploaded_names = list(st.session_state.documents.keys())

    if uploaded_names:
        st.divider()
        remove_name = st.selectbox(
            "Remove an uploaded document",
            options=uploaded_names,
        )

        if st.button("Remove selected document"):
            st.session_state.documents.pop(remove_name, None)
            st.cache_data.clear()
            st.success(f"Removed {remove_name}")
            st.rerun()


# ============================================================
# SETTINGS PAGE
# ============================================================
elif page == "⚙️ Settings":

    st.subheader("⚙️ Retrieval settings")

    st.markdown("""
    This application uses semantic vector search. Query embeddings
    are compared against document chunk embeddings using cosine
    similarity implemented by FAISS inner-product search on normalized
    vectors.
    """)

    st.markdown("### Current configuration")

    st.code("""
Embedding model:
sentence-transformers/all-MiniLM-L6-v2

Vector index:
FAISS IndexFlatIP

Embedding normalization:
Enabled

Chunk size:
900 characters

Chunk overlap:
150 characters

Default results:
Top 5 passages

Default similarity threshold:
0.20
""")

    st.warning(
        "The default threshold is a heuristic, not a calibrated "
        "relevance guarantee. Adjust it after evaluating queries "
        "against your actual policy documents."
    )

    st.markdown("### Runtime")
    st.write("Python application running under Streamlit.")
    st.write(f"Last interface refresh: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    st.markdown("### Limitations")
    st.markdown("""
    - Retrieved passages are not equivalent to verified answers.
    - The built-in HR policy content is illustrative demo data.
    - Scanned PDFs may require OCR before text can be extracted.
    - Uploaded documents are held in session memory and are not
      a persistent document database.
    - This app does not currently call an LLM to synthesize answers.
    """)


# ============================================================
# FOOTER
# ============================================================
st.divider()
st.markdown(
    '<div class="small-muted" style="text-align:center;">'
    'NEXUS HR KNOWLEDGE ASSISTANT · Semantic retrieval with FAISS · '
    'Verify policy details against the original document.'
    '</div>',
    unsafe_allow_html=True,
)